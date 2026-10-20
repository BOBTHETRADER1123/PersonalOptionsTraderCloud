import pytest
from io import BytesIO
from uuid import uuid4
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.database import Base, engine
from app.exporter import flush_exports
from app.main import app


@pytest.fixture
def client():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as test_client:
        login = test_client.post("/api/auth/login", json={"password": "paper-only"})
        assert login.status_code == 200
        test_client.headers.update({"X-CSRF-Token": login.json()["csrf_token"]})
        yield test_client
    flush_exports()
    Base.metadata.drop_all(engine)


def test_one_action_creates_independent_paper_trades_and_retries_are_idempotent(client):
    payload = {"action": "NEW_TRADE", "action_id": "action-test-1", "brokers": ["Dhan", "Kotak Neo", "5paisa"], "underlying": "NIFTY", "strategy": "Long Call", "legs": [{"underlying": "NIFTY", "expiry": "2030-01-31", "strike": 25000, "option_type": "CE", "side": "BUY", "lots": 1, "lot_size": 25, "entry_price": 100}]}
    first = client.post("/api/actions", json=payload, headers={"Idempotency-Key": "paper-key-1"})
    again = client.post("/api/actions", json=payload, headers={"Idempotency-Key": "paper-key-1"})
    assert first.status_code == 200
    assert len(first.json()["results"]) == 3
    assert again.json() == first.json()
    assert len(client.get("/api/state").json()["trades"]) == 3
    assert first.json()["mode"] == "PAPER"


def test_sell_to_close_calculates_realized_pnl_and_does_not_send_orders(client):
    created = client.post("/api/actions", json={"action": "NEW_TRADE", "brokers": ["Dhan"], "legs": [{"underlying": "NIFTY", "expiry": "2030-01-31", "strike": 25000, "option_type": "PE", "side": "SELL", "lots": 2, "lot_size": 25, "entry_price": 100}]})
    trade = created.json()["results"][0]
    exit_result = client.post("/api/actions", json={"action": "EXIT_STRATEGY", "trade_id": trade["id"], "exit_price": 70})
    assert exit_result.status_code == 200
    closed = exit_result.json()["results"][0]
    assert closed["status"] == "CLOSED"
    assert closed["realized_pnl"] == 1500
    assert client.get("/api/live-trading").json()["available"] is False


def test_limit_simulation_waits_for_qualifying_broker_quote_then_fills(client):
    created = client.post("/api/actions", json={"action": "NEW_TRADE", "order_type": "LIMIT", "brokers": ["Dhan"], "legs": [{"underlying": "NIFTY", "expiry": "2030-01-31", "strike": 25000, "option_type": "CE", "side": "BUY", "lots": 1, "lot_size": 25, "entry_price": 100}]})
    trade = created.json()["results"][0]
    assert trade["status"] == "PENDING"
    assert trade["legs"][0]["open_lots"] == 0
    quote = {"instrument_id": "security-1", "underlying": "NIFTY", "exchange": "NSE", "segment": "NSE_FNO", "symbol": trade["legs"][0]["symbol"], "expiry": "2030-01-31", "strike": 25000, "option_type": "CE", "broker": "Dhan", "ltp": 99, "sequence": f"seq-{uuid4()}"}
    response = client.post("/api/quotes", json=quote)
    assert response.status_code == 200
    state = client.get("/api/state").json()
    filled = next(x for x in state["trades"] if x["id"] == trade["id"])
    assert filled["status"] == "OPEN"
    assert filled["legs"][0]["open_lots"] == 1
    assert filled["legs"][0]["avg_entry"] == 99
    exported = client.get(f"/api/trades/{trade['id']}/export")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["Trade Summary", "P&L Timeline", "Trade Events", "Strategy Legs"]
    assert workbook["P&L Timeline"].max_row == 2


def test_pending_limit_simulation_can_be_cancelled(client):
    created = client.post("/api/actions", json={"action": "NEW_TRADE", "order_type": "LIMIT", "brokers": ["Dhan"], "legs": [{"underlying": "NIFTY", "expiry": "2030-01-31", "strike": 25000, "option_type": "PE", "side": "SELL", "lots": 1, "lot_size": 25, "entry_price": 80}]})
    trade = created.json()["results"][0]
    result = client.post("/api/actions", json={"action": "CANCEL_PENDING", "trade_id": trade["id"]})
    assert result.status_code == 200
    assert result.json()["results"][0]["status"] == "CANCELLED"


def test_csrf_required_for_mutations(client):
    client.headers.pop("X-CSRF-Token", None)
    response = client.post("/api/actions", json={"action": "NEW_TRADE", "brokers": ["Dhan"]})
    assert response.status_code == 403
