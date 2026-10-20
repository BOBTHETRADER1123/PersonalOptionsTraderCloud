from concurrent.futures import Future, ThreadPoolExecutor, wait
from pathlib import Path
from threading import Lock
import logging
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
from .config import settings
from .domain import payoff_summary
from .models import Leg, PnLSnapshot, Trade, TradeEvent

_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="xlsx-export")
_pending: set[Future] = set()
_pending_lock = Lock()
logger = logging.getLogger(__name__)


def workbook_path(trade: Trade) -> Path:
    stamp = trade.created_at
    if stamp.tzinfo is None:
        from datetime import timezone
        stamp = stamp.replace(tzinfo=timezone.utc)
    from zoneinfo import ZoneInfo
    stamp = stamp.astimezone(ZoneInfo("Asia/Kolkata"))
    year, month = stamp.strftime("%Y"), stamp.strftime("%m")
    safe_broker = {"Kotak Neo": "KotakNeo", "5paisa": "5paisa", "Dhan": "Dhan"}.get(trade.broker, trade.broker)
    folder = settings.data_path / safe_broker / year / month / f"Trade_{trade.id}"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"Trade_{trade.id}.xlsx"


def write_trade_workbook(db: Session, trade_id: str) -> Path:
    trade = db.get(Trade, trade_id)
    if not trade:
        raise ValueError("Trade not found")
    path = workbook_path(trade)
    wb = Workbook()
    summary = wb.active
    summary.title = "Trade Summary"
    legs = list(trade.legs)
    events = list(db.query(TradeEvent).filter_by(trade_id=trade_id).order_by(TradeEvent.occurred_at).all())
    unrealized = sum(x.direction * (x.mark - x.avg_entry) * x.open_lots * x.lot_size for x in legs)
    realized = sum(x.realized_pnl for x in legs)
    by_leg = {x.id: x for x in legs}
    premium_flow = 0.0
    for event in events:
        leg = by_leg.get(event.leg_id)
        details = event.details or {}
        price, lots = details.get("fill_price"), details.get("lots")
        if not leg or price is None:
            continue
        quantity = int(lots or 0) * leg.lot_size
        if event.event_type in {"ENTRY", "ADDED_LOTS", "LIMIT_FILLED"}:
            premium_flow -= leg.direction * float(price) * quantity
        elif event.event_type in {"PARTIAL_EXIT", "EXIT"}:
            premium_flow += leg.direction * float(price) * quantity
    risk_legs = [{"side": "BUY" if x.direction == 1 else "SELL", "option_type": x.option_type, "strike": x.strike, "entry_price": x.avg_entry, "lots": x.open_lots, "lot_size": x.lot_size} for x in legs if x.open_lots]
    risk = payoff_summary(risk_legs)
    max_profit = None if risk["max_profit"] is None else realized + risk["max_profit"]
    max_loss = None if risk["max_loss"] is None else realized + risk["max_loss"]
    summary.append(["Field", "Value"])
    rows = [
        ("Trade ID", trade.id), ("Strategy ID", trade.id), ("Broker", trade.broker),
        ("Underlying", trade.underlying), ("Strategy", trade.strategy), ("Status", trade.status),
        ("Created (UTC)", str(trade.created_at)), ("Closed (UTC)", str(trade.closed_at or "")),
        ("Gross realized P&L", realized), ("Gross unrealized P&L", unrealized),
        ("Charges", sum(x.charges for x in legs)),
        ("Net P&L", realized + unrealized - sum(x.charges for x in legs)),
        ("Maximum recorded profit (net P&L)", trade.max_profit_pnl if trade.max_profit_pnl is not None else "No quote recorded"),
        ("Maximum profit timestamp (UTC)", str(trade.max_profit_at or "")),
        ("Maximum recorded loss (net P&L)", trade.max_loss_pnl if trade.max_loss_pnl is not None else "No quote recorded"),
        ("Maximum loss timestamp (UTC)", str(trade.max_loss_at or "")),
        ("Net premium cash flow", premium_flow),
        ("Maximum profit estimate", "Uncapped" if max_profit is None else max_profit),
        ("Maximum loss estimate", "Unbounded" if max_loss is None else max_loss),
        ("Breakeven prices", ", ".join(str(x) for x in risk["breakevens"])),
        ("Margin estimate", "Unavailable: no broker margin response"),
        ("Hedge benefit estimate", "Unavailable: no broker margin response"),
    ]
    for row in rows:
        summary.append(list(row))

    timeline = wb.create_sheet("P&L Timeline")
    timeline.append(["Event timestamp UTC", "Instrument", "Broker", "LTP", "Bid", "Ask", "Open lots", "Average entry", "Unrealized P&L", "Realized P&L", "Charges", "Net P&L", "Feed age ms", "Event ID"])
    for snap in db.query(PnLSnapshot).filter_by(trade_id=trade_id).order_by(PnLSnapshot.event_at, PnLSnapshot.id).all():
        timeline.append([str(snap.event_at), snap.instrument, snap.source_broker, snap.ltp, snap.bid, snap.ask, snap.open_lots, snap.avg_entry, snap.unrealized_pnl, snap.realized_pnl, snap.charges, snap.net_pnl, snap.feed_age_ms, snap.sequence])

    event_sheet = wb.create_sheet("Trade Events")
    event_sheet.append(["Timestamp UTC", "Event", "Leg ID", "Action ID", "Details"])
    for event in events:
        event_sheet.append([str(event.occurred_at), event.event_type, event.leg_id, event.action_id, str(event.details)])

    leg_sheet = wb.create_sheet("Strategy Legs")
    leg_sheet.append(["Leg ID", "Broker", "Underlying", "Expiry", "Strike", "CE/PE", "Side", "Symbol", "Lot size", "Open lots", "Avg entry", "Current price", "Realized P&L", "Unrealized P&L", "Charges", "Margin estimate", "Hedge benefit estimate"])
    for leg in legs:
        leg_sheet.append([leg.id, leg.broker, leg.underlying, leg.expiry, leg.strike, leg.option_type, "BUY" if leg.direction == 1 else "SELL", leg.symbol, leg.lot_size, leg.open_lots, leg.avg_entry, leg.mark, leg.realized_pnl, leg.direction * (leg.mark - leg.avg_entry) * leg.open_lots * leg.lot_size, leg.charges, "Unavailable", "Unavailable"])

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else "A1"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="14243A")
        for col in ws.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in list(col)[:200]) + 2, 12), 42)
            ws.column_dimensions[get_column_letter(col[0].column)].width = width

    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    try:
        tmp.replace(path)
        return path
    except PermissionError:
        # Excel may hold the destination open on Windows. Keep this report as a separate recovery copy.
        from datetime import datetime, timezone
        alternate = path.with_name(f"{path.stem}.pending-{datetime.now(timezone.utc).strftime('%H%M%S')}.xlsx")
        tmp.replace(alternate)
        logger.warning("Workbook destination is locked; wrote a separate recovery copy: %s", alternate)
        return alternate


def write_async(trade_id: str, database_session_factory) -> None:
    def task():
        db = database_session_factory()
        try:
            write_trade_workbook(db, trade_id)
        finally:
            db.close()

    future = _pool.submit(task)
    with _pending_lock:
        _pending.add(future)

    def finished(done: Future):
        with _pending_lock:
            _pending.discard(done)
        error = done.exception()
        if error:
            logger.error("Workbook export failed for trade %s", trade_id, exc_info=(type(error), error, error.__traceback__))

    future.add_done_callback(finished)


def flush_exports(timeout: float = 10.0) -> bool:
    with _pending_lock:
        pending = list(_pending)
    if not pending:
        return True
    _, unfinished = wait(pending, timeout=timeout)
    return not unfinished
